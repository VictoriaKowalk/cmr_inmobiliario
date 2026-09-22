from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "documentacion"
OUTPUT = OUTPUT_DIR / "Manual_breve_panel_inmobiliario.docx"

BLUE = "2563A8"
DARK = "17324D"
MUTED = "667085"
LIGHT = "EAF2F8"
VERY_LIGHT = "F7F9FC"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, 9, MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_char_2)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, 27, DARK, True)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(18)
        run2 = p2.add_run(subtitle)
        set_font(run2, 13, MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    return p, run


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, 11, DARK, True)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest, 11, "1F2937")
    else:
        run = p.add_run(text)
        set_font(run, 11, "1F2937")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_font(run, 11, "1F2937")
    return p


def add_step(doc, number, title, detail):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [720, 8640])
    left, right = table.rows[0].cells
    set_cell_shading(left, BLUE)
    set_cell_shading(right, VERY_LIGHT)
    p_left = left.paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_left = p_left.add_run(str(number))
    set_font(r_left, 14, WHITE, True)
    p_right = right.paragraphs[0]
    p_right.paragraph_format.space_after = Pt(2)
    r_title = p_right.add_run(title)
    set_font(r_title, 11, DARK, True)
    p_detail = right.add_paragraph()
    p_detail.paragraph_format.space_after = Pt(0)
    r_detail = p_detail.add_run(detail)
    set_font(r_detail, 10.5, "344054")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{title}: ")
    set_font(r1, 10.5, DARK, True)
    r2 = p.add_run(text)
    set_font(r2, 10.5, DARK)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("1F2937")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 14, 7),
        ("Heading 3", 12, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if name != "Heading 3" else DARK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("GUÍA DE USO  |  PANEL INMOBILIARIO")
    set_font(hr, 8.5, MUTED, True)
    add_page_number(section.footer.paragraphs[0])

    add_title(doc, "Manual breve del panel inmobiliario", "Guía práctica para la administración diaria")
    add_note(
        doc,
        "Objetivo",
        "El panel permite mantener actualizadas las propiedades y atender desde un único lugar las consultas y solicitudes de tasación recibidas.",
    )

    add_heading(doc, "1. Ingreso al panel")
    add_body(doc, "Acceda a la dirección proporcionada por el administrador e ingrese su correo electrónico y contraseña.")
    add_bullet(doc, "La sesión es personal: no comparta su usuario ni su contraseña.")
    add_bullet(doc, "Para finalizar, utilice la opción “Salir” del menú.")
    add_bullet(doc, "Desde “Mi cuenta” puede cambiar su contraseña.")

    add_heading(doc, "2. Pantalla principal")
    add_body(
        doc,
        "El Dashboard resume la actividad del sistema y permite detectar rápidamente propiedades, consultas y tareas que requieren atención.",
    )
    add_body(
        doc,
        "El menú principal organiza las funciones en Propiedades, Contactos, Tipos de propiedad, Ubicaciones, Características, Usuarios y Cuenta.",
    )

    add_heading(doc, "3. Publicar o modificar una propiedad")
    add_step(doc, 1, "Crear la ficha", "Ingrese en Propiedades y seleccione “Nueva propiedad”. Complete la información general y la ubicación.")
    add_step(doc, 2, "Definir la operación", "Indique si la propiedad está en venta, alquiler o alquiler temporal, junto con precio, moneda y estado.")
    add_step(doc, 3, "Agregar contenido", "Cargue imágenes, elija una portada, ordene la galería y agregue videos de YouTube si corresponde.")
    add_step(doc, 4, "Revisar y guardar", "Verifique descripción, características y datos comerciales antes de confirmar.")

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_title(doc, "Gestión cotidiana", "Las tareas más frecuentes dentro del panel")

    add_heading(doc, "4. Propiedades")
    add_bullet(doc, "Buscar y filtrar propiedades existentes.")
    add_bullet(doc, "Editar datos generales, ubicación, características y operaciones.")
    add_bullet(doc, "Marcar una propiedad como destacada.")
    add_bullet(doc, "Cambiar rápidamente el estado comercial de una operación.")
    add_bullet(doc, "Eliminar una propiedad y recuperarla posteriormente desde la papelera.")
    add_note(
        doc,
        "Importante",
        "Antes de eliminar una propiedad, confirme que ya no deba aparecer publicada. La papelera permite restaurarla mientras el registro permanezca disponible.",
    )

    add_heading(doc, "5. Contactos, consultas y tasaciones")
    add_body(
        doc,
        "La sección Contactos reúne los mensajes enviados desde los formularios públicos del sitio. Cada registro conserva los datos de la persona, el tipo de solicitud y su seguimiento.",
    )
    add_bullet(doc, "Consultas generales: mensajes enviados desde el formulario de contacto.")
    add_bullet(doc, "Consultas por propiedad: solicitudes asociadas a una publicación específica.")
    add_bullet(doc, "Tasaciones: pedidos de valoración de una propiedad.")
    add_body(
        doc,
        "Al atender un contacto, actualice su estado de seguimiento. Esto permite distinguir los ingresos nuevos de los casos ya gestionados.",
    )

    add_heading(doc, "6. Catálogos auxiliares")
    add_body(doc, "Estos módulos alimentan las opciones disponibles al crear o editar una propiedad:")
    add_bullet(doc, "Tipos de propiedad: casa, departamento, terreno u otras categorías.")
    add_bullet(doc, "Ubicaciones: estructura geográfica utilizada para localizar los inmuebles.")
    add_bullet(doc, "Características: servicios, ambientes, amenities y demás atributos.")
    add_body(
        doc,
        "Los elementos pueden activarse o desactivarse. Conviene evitar categorías duplicadas y utilizar nombres claros y consistentes.",
    )

    add_heading(doc, "7. Usuarios administradores")
    add_body(
        doc,
        "Los usuarios habilitados pueden ingresar al panel. Desde este módulo es posible crear administradores, modificar sus datos o desactivar accesos.",
    )
    add_note(
        doc,
        "Seguridad",
        "Desactive inmediatamente las cuentas de personas que ya no deban operar el sistema. Cada persona debe usar su propio acceso.",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_title(doc, "Buenas prácticas", "Recomendaciones para mantener el panel ordenado y seguro")

    add_heading(doc, "8. Revisión antes de publicar")
    add_bullet(doc, "Título y descripción claros, sin errores de escritura.")
    add_bullet(doc, "Ubicación y tipo de propiedad correctos.")
    add_bullet(doc, "Precio, moneda y modalidad comercial actualizados.")
    add_bullet(doc, "Imagen de portada de buena calidad y galería ordenada.")
    add_bullet(doc, "Estado de disponibilidad correcto.")
    add_bullet(doc, "Datos sensibles del propietario fuera de la descripción pública.")

    add_heading(doc, "9. Mantenimiento recomendado")
    add_bullet(doc, "Revisar diariamente los contactos nuevos.")
    add_bullet(doc, "Actualizar el estado de propiedades vendidas, alquiladas o pausadas.")
    add_bullet(doc, "Evitar cargar dos veces la misma propiedad.")
    add_bullet(doc, "Usar contraseñas extensas y únicas.")
    add_bullet(doc, "Cerrar la sesión al utilizar equipos compartidos.")

    add_heading(doc, "10. Si algo no funciona")
    add_body(doc, "Antes de solicitar asistencia:")
    add_bullet(doc, "Actualice la página y vuelva a intentar la operación.")
    add_bullet(doc, "Compruebe que los campos obligatorios estén completos.")
    add_bullet(doc, "Verifique el tamaño y formato de imágenes o videos.")
    add_bullet(doc, "Anote qué estaba haciendo y, si es posible, tome una captura del mensaje de error.")

    add_note(
        doc,
        "Soporte",
        "Para recuperar accesos, resolver errores o realizar cambios técnicos, comuníquese con la persona responsable del mantenimiento del sistema.",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Panel inmobiliario · Manual de uso para el cliente")
    set_font(r, 9, MUTED, True)

    doc.core_properties.title = "Manual breve del panel inmobiliario"
    doc.core_properties.subject = "Guía de uso para el cliente"
    doc.core_properties.author = "Hemisferio Sur"
    doc.core_properties.keywords = "panel inmobiliario, manual, propiedades, contactos"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
